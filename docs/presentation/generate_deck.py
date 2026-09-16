#!/usr/bin/env python3
"""Generate a Google-ready .pptx deck for ARTEMIS.

Authored by Bassam S Faraj Jr., Product Evangelist & Technical Copywriter.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Theme ---
PRIMARY = RGBColor(0x1A, 0x73, 0xE8)  # Google-blue-ish
DARK = RGBColor(0x20, 0x20, 0x20)
LIGHT = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x00, 0x96, 0x88)
GRAY = RGBColor(0x66, 0x66, 0x66)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def set_slide_size(prs):
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT


def add_background(slide, fill_color=LIGHT):
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT
    )
    background.fill.solid()
    background.fill.fore_color.rgb = fill_color
    background.line.fill.background()
    # Send to back
    slide.shapes._spTree.remove(background._element)
    slide.shapes._spTree.insert(2, background._element)


def add_textbox(slide, left, top, width, height, text, font_size=18, bold=False,
                color=DARK, align=PP_ALIGN.LEFT, font_name="Google Sans", italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    run.font.italic = italic
    return box


def add_bullet_list(slide, left, top, width, height, items, font_size=16,
                    color=DARK, font_name="Google Sans"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {item}"
        p.level = 0
        p.space_after = Pt(10)
        run = p.runs[0]
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        run.font.name = font_name
    return box


def add_footer(slide, text):
    add_textbox(
        slide,
        left=Inches(0.5),
        top=SLIDE_HEIGHT - Inches(0.45),
        width=Inches(12),
        height=Inches(0.3),
        text=text,
        font_size=9,
        color=RGBColor(0x66, 0x66, 0x66),
        align=PP_ALIGN.LEFT,
    )


def add_section_header(slide, title, subtitle=None):
    add_textbox(
        slide,
        left=Inches(0.75),
        top=Inches(0.6),
        width=Inches(11.8),
        height=Inches(1.0),
        text=title,
        font_size=40,
        bold=True,
        color=PRIMARY,
        align=PP_ALIGN.LEFT,
    )
    if subtitle:
        add_textbox(
            slide,
            left=Inches(0.75),
            top=Inches(1.5),
            width=Inches(11.8),
            height=Inches(0.6),
            text=subtitle,
            font_size=20,
            color=DARK,
            align=PP_ALIGN.LEFT,
        )


def build_deck():
    prs = Presentation()
    set_slide_size(prs)
    blank_layout = prs.slide_layouts[6]

    # --- Slide 1: Title ---
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    add_textbox(
        slide,
        left=Inches(0.75),
        top=Inches(2.2),
        width=Inches(11.8),
        height=Inches(1.4),
        text="ARTEMIS",
        font_size=64,
        bold=True,
        color=PRIMARY,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        left=Inches(0.75),
        top=Inches(3.6),
        width=Inches(11.8),
        height=Inches(0.8),
        text="Autonomous AI Assistant for Mobile Automation",
        font_size=28,
        color=DARK,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        left=Inches(0.75),
        top=Inches(4.5),
        width=Inches(11.8),
        height=Inches(0.6),
        text="Let AI assistants and test suites use real phones like a human.",
        font_size=20,
        italic=True,
        color=ACCENT,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        left=Inches(0.75),
        top=Inches(5.8),
        width=Inches(11.8),
        height=Inches(0.5),
        text="Bassam S Faraj Jr. · Product Evangelist & Technical Copywriter",
        font_size=14,
        color=RGBColor(0x66, 0x66, 0x66),
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "ARTEMIS Product Deck")

    # --- Slide 2: The Problem ---
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    add_section_header(slide, "The Problem", "Why mobile automation still breaks teams")
    add_bullet_list(
        slide,
        left=Inches(1.0),
        top=Inches(2.5),
        width=Inches(11.0),
        height=Inches(4.0),
        items=[
            "Brittle selectors break when buttons move or IDs change.",
            "Custom renderers — Canvas, Compose, Flutter — evade traditional tools.",
            "Cross-app journeys require manual hand-offs and glue scripts.",
            "Transient UI (toasts, auto-fading bars) is impossible to catch reliably.",
            "IDEs lack native device diagnostics: screenshots, logcat, and replay.",
        ],
        font_size=20,
    )
    add_footer(slide, "ARTEMIS Product Deck · Bassam S Faraj Jr.")

    # --- Slide 3: Product Lineup ---
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    add_section_header(slide, "Product Lineup", "Four interfaces, one autonomous engine")
    items = [
        "Web Console — visual test orchestration, live mirroring, and replay",
        "Developer CLI — CI/CD runs, benchmarks, and terminal-first workflows",
        "MCP Server — native IDE control in Antigravity, Claude Code, Windsurf, Codex",
        "Python SDK — embed into pytest or existing automation with typed outputs",
        "Flash & Pro profiles — choose speed or deep planning and verification",
    ]
    add_bullet_list(
        slide,
        left=Inches(1.0),
        top=Inches(2.4),
        width=Inches(11.0),
        height=Inches(4.2),
        items=items,
        font_size=19,
    )
    add_footer(slide, "ARTEMIS Product Deck · Bassam S Faraj Jr.")

    # --- Slide 4: Core Features ---
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    add_section_header(slide, "Core Features", "Multimodal perception and closed-loop execution")
    items = [
        "Cross-app automation across 20+ apps from natural language",
        "Multimodal targeting: accessibility XML, OCR, and visual grounding",
        "Dynamic-first, coordinate-fallback locator philosophy",
        "Reactive observe-think-act loop with asynchronous history summaries",
        "Click sequences for transient controls",
        "Shared history compression: screenshots fold into visual summaries",
        "Safety Net pre-execution checks for Pro actions",
    ]
    add_bullet_list(
        slide,
        left=Inches(1.0),
        top=Inches(2.4),
        width=Inches(11.0),
        height=Inches(4.2),
        items=items,
        font_size=18,
    )
    add_footer(slide, "ARTEMIS Product Deck · Bassam S Faraj Jr.")

    # --- Slide 5: Benefits ---
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    add_section_header(slide, "Benefits", "What engineering and QA teams gain")
    items = [
        "Write tests in plain English instead of XPath or coordinates",
        "Run against real devices for production-class confidence",
        "Debug inside your IDE with screenshots and logcat",
        "Plug structured Pydantic outputs directly into CI gates",
        "Compress long session history without losing recall",
        "Reduce maintenance as UI drifts — locators adapt first, fallback second",
    ]
    add_bullet_list(
        slide,
        left=Inches(1.0),
        top=Inches(2.5),
        width=Inches(11.0),
        height=Inches(4.0),
        items=items,
        font_size=19,
    )
    add_footer(slide, "ARTEMIS Product Deck · Bassam S Faraj Jr.")

    # --- Slide 6: Solutions ---
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    add_section_header(slide, "Solutions", "Proof points and integrations")
    items = [
        "AndroidWorld SOTA: 99%+ task completion across 100+ multi-step tasks",
        "One-click MCP and agent-rule installation for major AI IDEs",
        "Flash Profile: 3–5 seconds per step for routine deterministic work",
        "Pro Profile: 15–40 seconds per step with Planner, Operator, and Checker",
        "Fast-action bursts defeat model turn latency on transient UI",
        "Execution incidents stay in context until recovery succeeds",
    ]
    add_bullet_list(
        slide,
        left=Inches(1.0),
        top=Inches(2.5),
        width=Inches(11.0),
        height=Inches(4.0),
        items=items,
        font_size=19,
    )
    add_footer(slide, "ARTEMIS Product Deck · Bassam S Faraj Jr.")

    # --- Slide 7: Evolution Path ---
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    add_section_header(slide, "Evolution Path", "Where ARTEMIS is headed")
    items = [
        "Android Studio Integration — in-editor debugging and test recording",
        "iOS Platform Expansion — real devices and simulators",
        "On-Device Lightweight VLMs — privacy-first local execution",
        "Real-time Duplex Voice Interaction — hands-free task dispatch",
    ]
    add_bullet_list(
        slide,
        left=Inches(1.0),
        top=Inches(2.8),
        width=Inches(11.0),
        height=Inches(3.4),
        items=items,
        font_size=22,
    )
    add_footer(slide, "ARTEMIS Product Deck · Bassam S Faraj Jr.")

    # --- Slide 8: Overhead Impact ---
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    add_section_header(slide, "Overhead Impact", "Doing more with less manual cost")
    items = [
        "Flash cuts routine task latency to ~3–5 seconds per step",
        "Pro compresses history instead of truncating it — long tasks stay coherent",
        "Closed-loop recovery reduces flaky reruns and manual intervention",
        "Dynamic-first locators lower selector-maintenance burden",
        "One natural-language prompt replaces many lines of glue code",
    ]
    add_bullet_list(
        slide,
        left=Inches(1.0),
        top=Inches(2.6),
        width=Inches(11.0),
        height=Inches(3.8),
        items=items,
        font_size=20,
    )
    add_footer(slide, "ARTEMIS Product Deck · Bassam S Faraj Jr.")

    # --- Slide 9: Call to Action ---
    slide = prs.slides.add_slide(blank_layout)
    add_background(slide)
    add_section_header(slide, "Next Steps", "Start automating in minutes")
    items = [
        "Clone the repo and run ./start.sh (macOS/Linux) or start.bat (Windows)",
        "Connect an Android device or emulator with USB debugging enabled",
        "Try: uv run artemis run \"Open Settings, find Battery, and report the level\" --profile flash",
        "Install MCP configs: uv run artemis mcp --install all",
        "Read the Mobile Testing Mindset at mcp_server/rules.md",
    ]
    add_bullet_list(
        slide,
        left=Inches(1.0),
        top=Inches(2.6),
        width=Inches(11.0),
        height=Inches(3.8),
        items=items,
        font_size=19,
    )
    add_textbox(
        slide,
        left=Inches(0.75),
        top=Inches(6.3),
        width=Inches(11.8),
        height=Inches(0.5),
        text="github.com/google/artemis",
        font_size=18,
        bold=True,
        color=PRIMARY,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "ARTEMIS Product Deck · Bassam S Faraj Jr.")

    return prs


if __name__ == "__main__":
    deck = build_deck()
    out_path = "artemis_product_deck.pptx"
    deck.save(out_path)
    print(f"Saved presentation to {out_path}")
